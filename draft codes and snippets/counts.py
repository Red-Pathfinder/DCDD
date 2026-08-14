import pandas as pd
from pathlib import Path
from openpyxl.styles import PatternFill

# ─── PATHS ──────────────────────────────────────────────────────────────────────
FLUX_CSV = Path(r"C:\Study\Coding\Research\Drift\data_processed\01_flux_attributes\attributes.csv")
SD35_CSV = Path(r"C:\Study\Coding\Research\Drift\data_processed\02_sd35_attributes\attributes.csv")
OUTPUT_EXCEL = Path(r"C:\Study\Coding\Research\draft codes and snippets\Raw_Counts.xlsx")

# ─── ALL ATTRIBUTES ─────────────────────────────────────────────────────────────
VISUAL_ATTRIBUTE_COLS = [
    "gender_presentation", "perceived_race", "skin_tone", "perceived_age",
    "sentiment", "environment_type", "professional_attire", "gaze_direction",
    "physical_framing", "activity_engagement", "social_integration"
]

# ─── MORE SATURATED EXCEL COLORS ────────────────────────────────────────────────
FLUX_COLOR = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")  # Richer Light Blue
SD35_COLOR = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")  # Richer Medium Gray
TOTAL_COLOR = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid") # Richer Warm Yellow

def main():
    print("Loading datasets...")
    
    try:
        df_flux = pd.read_csv(FLUX_CSV)
        df_sd35 = pd.read_csv(SD35_CSV)
        df = pd.concat([df_flux, df_sd35], ignore_index=True)
        print(f"✓ Combined {len(df_flux)} FLUX rows and {len(df_sd35)} SD 3.5 rows.")
    except Exception as e:
        print(f"[!] Error loading CSV files.\nError: {e}")
        return

    # Clean up model names
    df['model'] = df['model'].replace({
        'fal-ai/stable-diffusion-v35-large': 'SD 3.5 Large',
        'fal-ai/flux/dev': 'FLUX.1 Dev'
    })

    # ─── FIX FOR NEUTRAL BASELINE ───
    # Replace NaN or empty strings in the disability column so Pandas doesn't drop them
    df['disability'] = df['disability'].fillna('None (Baseline)')
    df['disability'] = df['disability'].replace({'': 'None (Baseline)', ' ': 'None (Baseline)'})

    print("\nGenerating color-coded Excel file...")
    
    with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
        for col in VISUAL_ATTRIBUTE_COLS:
            if col in df.columns:
                
                counts = pd.crosstab(
                    index=[df['model'], df['image_type'], df['disability']], 
                    columns=df[col],
                    margins=True, 
                    margins_name="Total"
                )
                
                sheet_name = col[:31] 
                counts.to_excel(writer, sheet_name=sheet_name)
                
                worksheet = writer.sheets[sheet_name]
                
                # 1. Auto-adjust column widths
                for excel_col in worksheet.columns:
                    max_length = 0
                    column = excel_col[0].column_letter 
                    for cell in excel_col:
                        try: 
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    worksheet.column_dimensions[column].width = (max_length + 2)

                # 2. Apply Row Colors
                current_model = None
                
                for row in range(2, worksheet.max_row + 1):
                    cell_val = str(worksheet.cell(row=row, column=1).value)
                    
                    if "FLUX" in cell_val:
                        current_model = "FLUX"
                    elif "SD 3.5" in cell_val:
                        current_model = "SD35"
                    elif "Total" in cell_val or "All" in cell_val:
                        current_model = "TOTAL"

                    fill_to_use = None
                    if current_model == "FLUX":
                        fill_to_use = FLUX_COLOR
                    elif current_model == "SD35":
                        fill_to_use = SD35_COLOR
                    elif current_model == "TOTAL":
                        fill_to_use = TOTAL_COLOR

                    if fill_to_use:
                        for col_idx in range(1, worksheet.max_column + 1):
                            worksheet.cell(row=row, column=col_idx).fill = fill_to_use

                print(f"  ✓ Formatted & colored sheet: {sheet_name}")

    print(f"\n✨ DONE! Beautiful Excel file saved to: {OUTPUT_EXCEL.name}")

if __name__ == "__main__":
    main()