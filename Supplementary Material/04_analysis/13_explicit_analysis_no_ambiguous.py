"""
Purpose:
Conducts statistical analysis exclusively on records with no ambiguous labels.

Paper Sections:
Section 4.1 (Robustness Check)

Input Files:
05_data/attributes_v2_flux.csv, 05_data/attributes_v2_sd35.csv, 05_data/nanobanana_attributes_final.csv

Output Files:
results/DCDD_Explicit_Analysis_No_Ambiguous.xlsx

Dependencies:
pandas, numpy, scipy, statsmodels, openpyxl

Approximate Runtime:
< 1 minute
"""

import pandas as pd
import numpy as np
from pathlib import Path
from scipy.stats import chi2_contingency
from statsmodels.stats.multitest import multipletests
from openpyxl.styles import PatternFill, Font

# Dynamic relative path to supplementary root
BASE_DIR = Path(__file__).resolve().parent.parent


# ─── PATHS ──────────────────────────────────────────────────────────────────────
FLUX_CSV   = BASE_DIR / "05_data" / "attributes_v2_flux.csv"
SD35_CSV   = BASE_DIR / "05_data" / "attributes_v2_sd35.csv"
GEMINI_CSV = BASE_DIR / "05_data" / "nanobanana_attributes_final.csv"

OUTPUT_EXCEL = BASE_DIR / "results" / "DCDD_Explicit_Analysis_No_Ambiguous.xlsx"
OUTPUT_EXCEL.parent.mkdir(parents=True, exist_ok=True)

# ─── EXCLUSION TARGETS ──────────────────────────────────────────────────────────
# Any category containing these words will be dropped before statistical testing
EXCLUSION_KEYWORDS = ['ambiguous', 'unclear', 'indeterminate', 'safety refusal']

VISUAL_ATTRIBUTE_COLS = [
    "image_style", "gender_presentation", "perceived_race", "skin_tone",
    "sentiment", "gaze_direction", "physical_framing", "activity_engagement",
    "social_integration"
]

# ─── EXCEL COLORS (SATURATED FOR READABILITY) ───────────────────────────────────
FLUX_COLOR   = PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid")  
SD35_COLOR   = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")  
GEMINI_COLOR = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")  

LARGE_BIAS_COLOR    = PatternFill(start_color="EA9999", end_color="EA9999", fill_type="solid") 
MODERATE_BIAS_COLOR = PatternFill(start_color="F9CB9C", end_color="F9CB9C", fill_type="solid") 
SIG_TEXT_COLOR      = Font(color="9C0006", bold=True) 

def calculate_cramers_v(chi2, n, min_dim):
    if n == 0 or min_dim <= 1: return 0.0
    return np.sqrt(chi2 / (n * (min_dim - 1)))

def get_bias_severity(v_stat):
    if pd.isna(v_stat) or v_stat == "N/A": return "N/A"
    v = float(v_stat)
    if v >= 0.30: return "Large Bias 🔴"
    if v >= 0.15: return "Moderate Bias 🟠"
    if v >= 0.05: return "Small Bias 🟡"
    return "Negligible"

def harmonize_gemini_data(df):
    """Aligns Gemini's metadata structure to match FLUX/SD3.5."""
    df = df.copy()
    df['model'] = 'Gemini 3.1 Flash'
    df['image_type'] = np.where(df['condition'].str.lower() == 'neutral', 'neutral', 'conditioned')
    df['disability'] = np.where(df['condition'].str.lower() == 'neutral', 'None (Neutral Baseline)', df['condition'])
    return df

def is_explicit(label):
    """Checks if a category label is explicit (not ambiguous or a refusal)."""
    if pd.isna(label): return False
    text = str(label).lower()
    return not any(keyword in text for keyword in EXCLUSION_KEYWORDS)

def main():
    print("[*] Loading and harmonizing all 3 datasets...")
    try:
        df_flux = pd.read_csv(FLUX_CSV)
        df_sd35 = pd.read_csv(SD35_CSV)
        df_gemini_raw = pd.read_csv(GEMINI_CSV)
        
        df_flux['model'] = 'FLUX.1 Dev'
        df_sd35['model'] = 'SD 3.5 Large'
        df_gemini = harmonize_gemini_data(df_gemini_raw)
        
        df = pd.concat([df_flux, df_sd35, df_gemini], ignore_index=True)
    except Exception as e:
        print(f"[!] Error loading CSV files.\nError: {e}")
        return

    df['disability'] = df['disability'].fillna('None (Neutral Baseline)')
    df['disability'] = df['disability'].replace({'': 'None (Neutral Baseline)', ' ': 'None (Neutral Baseline)'})
    
    excel_sheets = {}
    print("[*] Running Explicit-Only calculations (Ambiguity Removed)...")

    for col in VISUAL_ATTRIBUTE_COLS:
        if col not in df.columns: continue

        # Filter the dataframe to ONLY include rows where the current attribute is explicit
        explicit_mask = df[col].apply(is_explicit)
        df_explicit = df[explicit_mask].copy()

        # Get the remaining clear categories
        categories = sorted([str(x) for x in df_explicit[col].dropna().unique()])
        records = []
        p_values_to_correct = []
        p_value_indices = [] 

        for model in ['FLUX.1 Dev', 'SD 3.5 Large', 'Gemini 3.1 Flash']:
            model_df = df_explicit[df_explicit['model'] == model]
            
            neutral_df = model_df[model_df['image_type'] == 'neutral']
            if neutral_df.empty: continue
            
            neutral_counts = neutral_df[col].value_counts().reindex(categories, fill_value=0)
            
            base_row = {
                "Model": model,
                "Condition": "None (Neutral Baseline)",
                **neutral_counts.to_dict(),
                "Total (Explicit)": neutral_counts.sum(),
                "FDR p-value": "N/A",
                "Significant?": "N/A",
                "Cramér's V": "N/A",
                "Bias Severity": "N/A" 
            }
            records.append(base_row)

            conditions = [d for d in model_df['disability'].unique() if d != 'None (Neutral Baseline)']
            
            for condition in conditions:
                cond_df = model_df[(model_df['image_type'] == 'conditioned') & (model_df['disability'] == condition)]
                cond_counts = cond_df[col].value_counts().reindex(categories, fill_value=0)
                
                obs = np.array([neutral_counts.values, cond_counts.values])
                valid_cols = obs.sum(axis=0) > 0
                obs_clean = obs[:, valid_cols]

                p_val, v_stat = np.nan, np.nan
                
                if obs_clean.shape[1] > 1 and obs_clean.sum() > 0:
                    chi2, p_val, dof, expected = chi2_contingency(obs_clean)
                    n = obs_clean.sum()
                    min_dim = min(obs_clean.shape)
                    v_stat = calculate_cramers_v(chi2, n, min_dim)
                
                v_stat_rounded = round(v_stat, 3) if not np.isnan(v_stat) else "N/A"
                
                cond_row = {
                    "Model": model,
                    "Condition": condition,
                    **cond_counts.to_dict(),
                    "Total (Explicit)": cond_counts.sum(),
                    "FDR p-value": np.nan,
                    "Significant?": "No",
                    "Cramér's V": v_stat_rounded,
                    "Bias Severity": get_bias_severity(v_stat_rounded) 
                }
                
                records.append(cond_row)
                
                if not np.isnan(p_val):
                    p_values_to_correct.append(p_val)
                    p_value_indices.append(len(records) - 1)

        if p_values_to_correct:
            reject, pvals_corrected, _, _ = multipletests(p_values_to_correct, alpha=0.05, method='fdr_bh')
            for i, idx in enumerate(p_value_indices):
                records[idx]["FDR p-value"] = pvals_corrected[i]
                if reject[i]:
                    records[idx]["Significant?"] = "YES ✓"

        for r in records:
            if isinstance(r["FDR p-value"], float) and not np.isnan(r["FDR p-value"]):
                r["FDR p-value"] = f"{r['FDR p-value']:.4f}"
            if pd.isna(r["FDR p-value"]): r["FDR p-value"] = "N/A"

        excel_sheets[col[:31]] = pd.DataFrame(records)

    print("[*] Formatting and exporting to Excel...")
    with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
        for sheet_name, df_sheet in excel_sheets.items():
            df_sheet.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.sheets[sheet_name]
            
            for excel_col in worksheet.columns:
                max_length = 0
                column = excel_col[0].column_letter 
                for cell in excel_col:
                    try: 
                        if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
                    except: pass
                worksheet.column_dimensions[column].width = min(max_length + 2, 30)

            for row in range(2, worksheet.max_row + 1):
                model_val = str(worksheet.cell(row=row, column=1).value)
                severity_val = str(worksheet.cell(row=row, column=worksheet.max_column).value)
                
                if "FLUX" in model_val: fill_color = FLUX_COLOR
                elif "SD 3.5" in model_val: fill_color = SD35_COLOR
                else: fill_color = GEMINI_COLOR
                    
                for col_idx in range(1, worksheet.max_column + 1):
                    worksheet.cell(row=row, column=col_idx).fill = fill_color

                sev_col_idx = worksheet.max_column
                if "Large Bias" in severity_val:
                    worksheet.cell(row=row, column=sev_col_idx).fill = LARGE_BIAS_COLOR
                    worksheet.cell(row=row, column=sev_col_idx).font = SIG_TEXT_COLOR
                elif "Moderate Bias" in severity_val:
                    worksheet.cell(row=row, column=sev_col_idx).fill = MODERATE_BIAS_COLOR
                    worksheet.cell(row=row, column=sev_col_idx).font = Font(bold=True)

    print(f"\n✨ EXPLICIT ANALYSIS COMPLETE! Ready for Srikant Sir: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    main()