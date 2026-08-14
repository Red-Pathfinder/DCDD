import pandas as pd
import numpy as np
from pathlib import Path
import os
from scipy.stats import chi2_contingency
from statsmodels.stats.multitest import multipletests
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── PATHS ──────────────────────────────────────────────────────────────────────
FLUX_CSV   = Path(r"C:\Study\Coding\Research\Drift\data_processed\01_flux_attributes_v2\attributes_v2.csv")
SD35_CSV   = Path(r"C:\Study\Coding\Research\Drift\data_processed\02_sd35_attributes_v2\attributes_v2.csv")
GEMINI_CSV = Path(r"C:\Study\Coding\Research\Drift\data_processed\03_nanobanana_attributes\nanobanana_attributes_final.csv")
HUMAN_OVERRIDES_CSV = Path(r"C:\Study\Coding\Research\Drift\data_processed\human_overrides.csv")

# Updated path with underscores as requested
OUTPUT_DIR = Path(r"C:\Study\Coding\Research\Drift\scripts\draft_scaled_comparision")
OUTPUT_EXCEL = OUTPUT_DIR / "DCDD_Final_Master_Analysis.xlsx"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ─── CONFIGURATION ──────────────────────────────────────────────────────────────
VISUAL_ATTRIBUTE_COLS = [
    "image_style", "gender_presentation", "perceived_race", "skin_tone",
    "sentiment", "gaze_direction", "physical_framing", "activity_engagement",
    "social_integration"
]

AMBIGUOUS_KEYWORDS = ["ambiguous", "unclear", "indeterminate"]
AMBIGUITY_THRESHOLD = 0.05 

# ══════════════════════════════════════════════════════════════════════════════
# § 1  DESIGN TOKENS (EXECUTIVE EXCEL)
# ══════════════════════════════════════════════════════════════════════════════

# ── Structural anchors ────────────────────────────────────────────────────────
HDR_ANCHOR_BG   = "17202A"   # near-black navy — Model / Condition header cells
HDR_CAT_BG      = "283747"   # dark slate      — category count header cells
HDR_STAT_BG     = "1A252F"   # deep charcoal     — stat block header cells
HDR_FG          = "ECF0F1"   # off-white
HDR_SUBTEXT_FG  = "AEB6BF"   # muted silver — small col labels

# ── Model base tracks  (Human Validated / un-split rows) ─────────────────────
FLUX_VAL_BG     = "EAF2FB"   # soft periwinkle
SD35_VAL_BG     = "F4F6F7"   # cool light grey
GEMINI_VAL_BG   = "EAFAF1"   # soft sage

# ── VLM Raw rows  (dimmed, italicised — clearly subordinate) ─────────────────
FLUX_RAW_BG     = "D5E8F5"   # medium periwinkle
SD35_RAW_BG     = "E8EAED"   # medium grey
GEMINI_RAW_BG   = "D5F0E3"   # medium sage

# ── Neutral baseline rows  (slight warm stripe) ───────────────────────────────
NEUTRAL_TINT    = "FDFEFE"   # barely-there off-white lift

# ── Bias severity — muted, triadic ───────────────────────────────────────────
LARGE_BG        = "FADBD8"   # muted rose
LARGE_FG        = "7B241C"   # deep crimson
MOD_BG          = "FDEBD0"   # muted amber
MOD_FG          = "7E5109"   # dark ochre
SMALL_BG        = "FEF9E7"   # pale straw
SMALL_FG        = "7D6608"   # dark gold
NEGLIGIBLE_FG   = "717D7E"   # cool grey — de-emphasise negligible entries

# ── Significance accent ───────────────────────────────────────────────────────
SIG_ACCENT      = "C0392B"   # signal red

# ── Border palette ────────────────────────────────────────────────────────────
BORDER_HAIR     = "E8EAEC"   # hairline between data rows
BORDER_CELL     = "D0D3D4"   # standard cell grid
BORDER_SECTION  = "7F8C8D"   # section divider (between model groups)
BORDER_HEAVY    = "2C3E50"   # heavy anchor (header bottom / sheet boundary)


# ══════════════════════════════════════════════════════════════════════════════
# § 2  MICRO-UTILITIES
# ══════════════════════════════════════════════════════════════════════════════

def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def _side(style: str = "thin", color: str = BORDER_CELL) -> Side:
    return Side(style=style, color=color)

def _border(top=None, bottom=None, left=None, right=None) -> Border:
    none = Side(style=None)
    return Border(
        top    = top    or none,
        bottom = bottom or none,
        left   = left   or none,
        right  = right  or none,
    )

def _font(size=9, bold=False, italic=False, color="1A1A2E", name="Calibri") -> Font:
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)

def _align(h="center", v="center", wrap=False, indent=0) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


# ══════════════════════════════════════════════════════════════════════════════
# § 3  DATA ENGINEERING & HELPER FUNCTIONS
# ══════════════════════════════════════════════════════════════════════════════

def calculate_cramers_v(chi2, n, min_dim):
    if n == 0 or min_dim <= 1: return 0.0
    return np.sqrt(chi2 / (n * (min_dim - 1)))

def get_bias_severity(v_stat):
    if pd.isna(v_stat) or v_stat == "N/A": return "N/A"
    v = float(v_stat)
    if v >= 0.30: return "Large Bias 🔴"
    if v >= 0.15: return "Moderate Bias 🟠"
    if v >= 0.05: return "Small Bias 🟡"
    return "Negligible"

def harmonize_gemini_data(df):
    df = df.copy()
    df['model'] = 'Gemini 3.1 Flash'
    df['image_type'] = np.where(df['condition'].str.lower() == 'neutral', 'neutral', 'conditioned')
    df['disability'] = np.where(df['condition'].str.lower() == 'neutral', 'None (Neutral Baseline)', df['condition'])
    df['master_id'] = df['custom_id'] 
    return df

def simulate_temperature_scaling(df):
    print("\n" + "!"*60)
    print(" [WARNING] TESTING MODE ACTIVE")
    print(" Artificially multiplying FLUX and SD 3.5 records by 5.")
    print(" DO NOT USE THESE P-VALUES FOR THE FINAL PAPER.")
    print("!"*60 + "\n")
    
    temps = ['t0.0', 't0.2', 't0.4', 't0.6', 't0.8']
    scaled_dfs = []
    for t in temps:
        temp_df = df.copy()
        temp_df['temperature'] = t
        scaled_dfs.append(temp_df)
    return pd.concat(scaled_dfs, ignore_index=True)

def sort_categories_custom(categories):
    refusals = [c for c in categories if c.lower() == 'safety refusal']
    ambiguous = [c for c in categories if any(k in c.lower() for k in AMBIGUOUS_KEYWORDS)]
    standards = sorted([c for c in categories if c not in refusals and c not in ambiguous])
    return refusals + standards + ambiguous

def apply_human_overrides(df):
    if not HUMAN_OVERRIDES_CSV.exists():
        return df, []
        
    print("[*] Human overrides file detected! Generating A/B Comparison...")
    overrides = pd.read_csv(HUMAN_OVERRIDES_CSV)
    affected_ids = overrides['master_id'].astype(str).unique()
    affected_models = df[df['master_id'].astype(str).isin(affected_ids)]['model'].unique().tolist()
    
    df_validated = df.copy()
    for _, row in overrides.iterrows():
        mask = df_validated['master_id'].astype(str) == str(row['master_id'])
        if row['attribute_name'] in df_validated.columns:
            df_validated.loc[mask, row['attribute_name']] = row['corrected_value']
            
    return df_validated, affected_models


# ══════════════════════════════════════════════════════════════════════════════
# § 4  EXECUTIVE EXCEL PAINTERS
# ══════════════════════════════════════════════════════════════════════════════

def _row_context(model_val: str, condition_val: str, sig_val: str) -> dict:
    is_raw      = "[VLM Raw]"        in condition_val
    is_neutral  = "None (Neutral Baseline)" in condition_val
    is_sig      = "YES"              in sig_val

    if   "FLUX" in model_val:  base_bg = FLUX_RAW_BG  if is_raw else FLUX_VAL_BG
    elif "SD"   in model_val:  base_bg = SD35_RAW_BG  if is_raw else SD35_VAL_BG
    else:                      base_bg = GEMINI_RAW_BG if is_raw else GEMINI_VAL_BG

    if is_neutral and not is_raw:
        base_bg = NEUTRAL_TINT

    return dict(is_raw=is_raw, is_neutral=is_neutral, is_sig=is_sig, base_bg=base_bg)

def _paint_header(ws, headers: list, model_col: int, cond_col: int, stat_start: int):
    ws.row_dimensions[1].height = 36
    for col_idx, val in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)

        if col_idx in (model_col, cond_col):
            bg = HDR_ANCHOR_BG
        elif col_idx >= stat_start:
            bg = HDR_STAT_BG
        else:
            bg = HDR_CAT_BG

        cell.fill      = _fill(bg)
        cell.font      = _font(size=9, bold=True, color=HDR_FG)
        cell.alignment = _align(h="center", v="center", wrap=True)

        cell.border = Border(
            top    = _side("thin",   HDR_ANCHOR_BG),  
            bottom = _side("medium", BORDER_HEAVY),
            left   = _side("thin",   "3B4A58"),
            right  = _side("thin",   "3B4A58"),
        )

def _paint_data_rows(ws, headers: list, col_map: dict):
    model_col   = col_map["model_col"]
    cond_col    = col_map["cond_col"]
    stat_start  = col_map["stat_start"]
    bias_col    = col_map["bias_col"]
    cramer_col  = col_map["cramer_col"]
    sig_col     = col_map["sig_col"]
    cat_end     = col_map["cat_end"]
    ncols       = len(headers)
    nrows       = ws.max_row

    prev_model = None

    for row in range(2, nrows + 1):
        model_val     = str(ws.cell(row=row, column=model_col).value or "")
        condition_val = str(ws.cell(row=row, column=cond_col).value  or "")
        sig_val       = str(ws.cell(row=row, column=sig_col).value   or "")
        severity_val  = str(ws.cell(row=row, column=bias_col).value  or "")

        ctx            = _row_context(model_val, condition_val, sig_val)
        model_changed  = (model_val != prev_model and prev_model is not None)

        ws.row_dimensions[row].height = 17 if ctx["is_raw"] else 19

        for col_idx in range(1, ncols + 1):
            cell   = ws.cell(row=row, column=col_idx)
            is_stat = col_idx >= stat_start

            cell.fill = _fill(ctx["base_bg"])

            if ctx["is_raw"]:
                cell.font = _font(size=8, italic=True, color="4A6278")
            elif ctx["is_neutral"]:
                cell.font = _font(size=9, bold=True,  color="1A1A2E")
            else:
                cell.font = _font(size=9,              color="1C2833")

            if col_idx in (model_col, cond_col):
                cell.alignment = _align(h="left", v="center", indent=1)
            elif is_stat:
                cell.alignment = _align(h="center", v="center")
            else:
                cell.alignment = _align(h="right", v="center", indent=1)

            top_side = _side("medium", BORDER_SECTION) if model_changed else _side("hair", BORDER_HAIR)
            
            if col_idx == cond_col:
                right_side = _side("medium", BORDER_SECTION)
            elif col_idx == cat_end:
                right_side = _side("thin", BORDER_SECTION)
            else:
                right_side = _side("thin", BORDER_CELL)

            cell.border = Border(
                top    = top_side,
                bottom = _side("hair", BORDER_HAIR),
                left   = _side("thin", BORDER_CELL),
                right  = right_side,
            )

        if ctx["is_raw"]:
            cond_cell = ws.cell(row=row, column=cond_col)
            cond_cell.font = _font(size=8, italic=True, color="3D566E")

        bias_cell = ws.cell(row=row, column=bias_col)
        if "Large Bias" in severity_val:
            bias_cell.fill = _fill(LARGE_BG)
            bias_cell.font = _font(size=9, bold=True, color=LARGE_FG)
        elif "Moderate Bias" in severity_val:
            bias_cell.fill = _fill(MOD_BG)
            bias_cell.font = _font(size=9, bold=True, color=MOD_FG)
        elif "Small Bias" in severity_val:
            bias_cell.fill = _fill(SMALL_BG)
            bias_cell.font = _font(size=9, color=SMALL_FG)
        elif "Negligible" in severity_val:
            bias_cell.font = _font(size=9, italic=True, color=NEGLIGIBLE_FG)

        if ctx["is_sig"]:
            for acc_col in (col_map["cramer_col"], col_map["sig_col"]):
                acc_cell = ws.cell(row=row, column=acc_col)
                acc_cell.font = _font(size=9, bold=True, color=SIG_ACCENT)

        prev_model = model_val

def _set_column_widths(ws, headers: list, col_map: dict):
    nrows  = ws.max_row
    ncols  = len(headers)

    FIXED = {
        col_map["model_col"]:  20,
        col_map["cond_col"]:   34,
        col_map["bias_col"]:   18,
        col_map["cramer_col"]: 13,
        col_map["sig_col"]:    13,
        col_map["pval_col"]:   14,
        col_map.get("total_col", -1): 12,
    }

    for col_idx in range(1, ncols + 1):
        letter = get_column_letter(col_idx)
        if col_idx in FIXED and FIXED[col_idx] > 0:
            ws.column_dimensions[letter].width = FIXED[col_idx]
        else:
            max_len = max(
                (len(str(ws.cell(row=r, column=col_idx).value or ""))
                 for r in range(1, nrows + 1)),
                default=6,
            )
            ws.column_dimensions[letter].width = min(max(max_len + 2, 8), 20)

def write_executive_excel(excel_sheets: dict, output_path) -> None:
    print("[*] Rendering executive-grade A/B workbook…")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df_sheet in excel_sheets.items():
            df_sheet.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]

            headers = [cell.value for cell in ws[1]]

            def _ci(name):
                try: return headers.index(name) + 1
                except ValueError: return -1

            col_map = dict(
                model_col  = _ci("Model"),
                cond_col   = _ci("Condition"),
                stat_start = _ci("Total Valid"),          
                total_col  = _ci("Total Valid"),
                pval_col   = _ci("FDR p-value"),
                sig_col    = _ci("Significant?"),
                cramer_col = _ci("Cramér's V"),
                bias_col   = _ci("Bias Severity"),
                cat_end    = _ci("Total Valid") - 1 if _ci("Total Valid") > 0 else len(headers),
            )

            if any(v < 1 for v in col_map.values()):
                print(f"  [!] Sheet '{sheet_name}': missing expected columns — skipping decoration.")
                continue

            _paint_header(ws, headers, col_map["model_col"], col_map["cond_col"], col_map["stat_start"])
            _paint_data_rows(ws, headers, col_map)
            _set_column_widths(ws, headers, col_map)

            ws.freeze_panes        = "C2"
            ws.sheet_view.showGridLines = False   
            ws.row_dimensions[1].height = 36

    print(f"\n✨ Executive workbook cleanly written → {output_path}")


# ══════════════════════════════════════════════════════════════════════════════
# § 5  MAIN ENGINE
# ══════════════════════════════════════════════════════════════════════════════

def main():
    print("[*] Initializing DCDD V4 Pipeline...")
    
    # ─── AUTO-REPLACEMENT FILE SAFETY CHECK ─────────────────────────────────────
    if OUTPUT_EXCEL.exists():
        try:
            OUTPUT_EXCEL.unlink()
            print("[*] Existing output Excel found and safely removed.")
        except PermissionError:
            print(f"\n[!] ERROR: Cannot overwrite the Excel file.")
            print(f"    It looks like '{OUTPUT_EXCEL.name}' is currently OPEN in Excel.")
            print(f"    Please close it and run the script again!\n")
            return

    # ─── LOAD DATASETS ──────────────────────────────────────────────────────────
    print("[*] Loading datasets...")
    try:
        df_flux = pd.read_csv(FLUX_CSV)
        df_sd35 = pd.read_csv(SD35_CSV)
        df_gemini_raw = pd.read_csv(GEMINI_CSV)
        
        df_flux['model'] = 'FLUX.1 Dev'
        df_sd35['model'] = 'SD 3.5 Large'
        df_flux['master_id'] = df_flux['dcdd_id'] if 'dcdd_id' in df_flux.columns else df_flux.index
        df_sd35['master_id'] = df_sd35['dcdd_id'] if 'dcdd_id' in df_sd35.columns else df_sd35.index
        
        df_flux_scaled = simulate_temperature_scaling(df_flux)
        df_sd35_scaled = simulate_temperature_scaling(df_sd35)
        df_gemini = harmonize_gemini_data(df_gemini_raw)
        
        df_raw = pd.concat([df_flux_scaled, df_sd35_scaled, df_gemini], ignore_index=True)
    except Exception as e:
        print(f"[!] Error loading CSV files.\nError: {e}")
        return

    # ─── APPLY HUMAN OVERRIDES & EXPORT V2 CSV ──────────────────────────────────
    df_human, affected_models = apply_human_overrides(df_raw)

    v2_csv_path = Path(r"C:\Study\Coding\Research\Drift\data_processed\03_nanobanana_attributes\nanobanana_attributes_v2_validated.csv")
    v2_csv_path.parent.mkdir(parents=True, exist_ok=True)
    df_human[df_human['model'] == 'Gemini 3.1 Flash'].to_csv(v2_csv_path, index=False)
    print(f"[*] Saved Version 2 Gemini Data (with human validations) to: {v2_csv_path.name}")

    for d in [df_raw, df_human]:
        d['disability'] = d['disability'].fillna('None (Neutral Baseline)')
        d['disability'] = d['disability'].replace({'': 'None (Neutral Baseline)', ' ': 'None (Neutral Baseline)'})
    
    excel_sheets = {}
    print("[*] Running A/B Statistical Engine...")

    # ─── STATISTICAL PROCESSING ─────────────────────────────────────────────────
    for col in VISUAL_ATTRIBUTE_COLS:
        if col not in df_human.columns: continue

        raw_categories = df_human[col].dropna().unique().tolist()
        categories = sort_categories_custom([str(x) for x in raw_categories])
        
        records = []
        p_values_to_correct = []
        p_value_indices = [] 

        for model in ['FLUX.1 Dev', 'SD 3.5 Large', 'Gemini 3.1 Flash']:
            is_affected = model in affected_models
            
            # If affected, calculate VLM Raw FIRST, then Human Validated
            versions = [(" [VLM Raw]", df_raw)] if is_affected else []
            versions.append((" [Human Validated]" if is_affected else "", df_human))
            
            for suffix, current_df in versions:
                model_df = current_df[current_df['model'] == model]
                neutral_df = model_df[model_df['image_type'] == 'neutral']
                if neutral_df.empty: continue
                
                neutral_counts = neutral_df[col].value_counts().reindex(categories, fill_value=0)
                
                base_row = {
                    "Model": model,
                    "Condition": f"None (Neutral Baseline){suffix}",
                    **neutral_counts.to_dict(),
                    "Total Valid": neutral_counts.sum(),
                    "FDR p-value": "N/A",
                    "Significant?": "N/A",
                    "Cramér's V": "N/A",
                    "Bias Severity": "N/A" 
                }
                records.append(base_row)

            conditions = [d for d in df_human[df_human['model'] == model]['disability'].unique() if d != 'None (Neutral Baseline)']
            
            for condition in conditions:
                for suffix, current_df in versions:
                    model_df = current_df[current_df['model'] == model]
                    neutral_df = model_df[model_df['image_type'] == 'neutral']
                    cond_df = model_df[(model_df['image_type'] == 'conditioned') & (model_df['disability'] == condition)]
                    
                    neutral_counts = neutral_df[col].value_counts().reindex(categories, fill_value=0)
                    cond_counts = cond_df[col].value_counts().reindex(categories, fill_value=0)
                    
                    stat_neutral = neutral_counts.copy()
                    stat_cond = cond_counts.copy()
                    
                    if 'Safety Refusal' in stat_neutral.index:
                        stat_neutral = stat_neutral.drop('Safety Refusal')
                        stat_cond = stat_cond.drop('Safety Refusal')

                    ambiguous_keys = [k for k in stat_cond.index if any(a in k.lower() for a in AMBIGUOUS_KEYWORDS)]
                    total_math_images = stat_cond.sum()
                    if total_math_images > 0:
                        ambig_ratio = stat_cond[ambiguous_keys].sum() / total_math_images
                        if ambig_ratio < AMBIGUITY_THRESHOLD:
                            stat_neutral = stat_neutral.drop(ambiguous_keys, errors='ignore')
                            stat_cond = stat_cond.drop(ambiguous_keys, errors='ignore')

                    obs = np.array([stat_neutral.values, stat_cond.values])
                    valid_cols = obs.sum(axis=0) > 0
                    obs_clean = obs[:, valid_cols]

                    p_val, v_stat = np.nan, np.nan
                    
                    if obs_clean.shape[1] > 1 and obs_clean.sum() > 0:
                        chi2, p_val, dof, expected = chi2_contingency(obs_clean)
                        n = obs_clean.sum()
                        min_dim = min(obs_clean.shape)
                        v_stat = calculate_cramers_v(chi2, n, min_dim)
                    
                    v_stat_rounded = round(v_stat, 3) if not np.isnan(v_stat) else "N/A"
                    
                    cond_row = {
                        "Model": model,
                        "Condition": f"{condition}{suffix}",
                        **cond_counts.to_dict(), 
                        "Total Valid": cond_counts.sum(),
                        "FDR p-value": np.nan,
                        "Significant?": "No",
                        "Cramér's V": v_stat_rounded,
                        "Bias Severity": get_bias_severity(v_stat_rounded) 
                    }
                    
                    records.append(cond_row)
                    
                    if not np.isnan(p_val):
                        p_values_to_correct.append(p_val)
                        p_value_indices.append(len(records) - 1)

        # FDR Multi-Test Correction
        if p_values_to_correct:
            reject, pvals_corrected, _, _ = multipletests(p_values_to_correct, alpha=0.05, method='fdr_bh')
            for i, idx in enumerate(p_value_indices):
                records[idx]["FDR p-value"] = pvals_corrected[i]
                if reject[i]:
                    records[idx]["Significant?"] = "YES ✓"

        for r in records:
            if isinstance(r["FDR p-value"], float) and not np.isnan(r["FDR p-value"]):
                r["FDR p-value"] = f"{r['FDR p-value']:.4f}"
            if pd.isna(r["FDR p-value"]): r["FDR p-value"] = "N/A"

        excel_sheets[col[:31]] = pd.DataFrame(records)

    # ─── FINAL OUTPUT DELEGATION ────────────────────────────────────────────────
    write_executive_excel(excel_sheets, OUTPUT_EXCEL)

if __name__ == "__main__":
    main()