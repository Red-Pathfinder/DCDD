import pandas as pd
from pathlib import Path
from openpyxl.styles import PatternFill, Font

# ─── PATHS ──────────────────────────────────────────────────────────────────────
ERASURE_CSV = Path(r"C:\Study\Coding\Research\results\profession_erasure.csv")
OUTPUT_DIR = Path(r"C:\Study\Coding\Research\results")
OUTPUT_EXCEL = OUTPUT_DIR / "DCDD_Executive_Summary_Profession_Erasure.xlsx"

# ─── EXCEL COLORS ───────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
SEVERE_DROP = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Red for massive drop
MODERATE_DROP = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid") # Orange for moderate

def main():
    print("Loading Profession Erasure data...")
    try:
        df = pd.read_csv(ERASURE_CSV)
    except FileNotFoundError:
        print(f"[!] Error: Could not find {ERASURE_CSV}")
        return

    # Clean up model names to match our other sheets
    df['model'] = df['model'].replace({'FLUX': 'FLUX.1 Dev', 'SD35': 'SD 3.5 Large'})
    
    # Ensure YES/NO is standard
    df['profession_depicted'] = df['profession_depicted'].str.upper().str.strip()

    print("Calculating Erasure Drop-offs...")
    results = []

    for model in ['FLUX.1 Dev', 'SD 3.5 Large']:
        model_df = df[df['model'] == model]
        if model_df.empty: continue

        # Calculate Neutral Baseline Adherence
        neutral_df = model_df[model_df['image_type'] == 'neutral']
        if not neutral_df.empty:
            neutral_yes = len(neutral_df[neutral_df['profession_depicted'] == 'YES'])
            baseline_adherence = (neutral_yes / len(neutral_df)) * 100
        else:
            baseline_adherence = 100.0 # Fallback

        # Calculate Adherence for EACH specific disability
        conditions = [d for d in model_df['disability'].unique() if pd.notna(d) and str(d).strip() != '']
        
        for condition in conditions:
            cond_df = model_df[(model_df['image_type'] == 'conditioned') & (model_df['disability'] == condition)]
            if cond_df.empty: continue
            
            cond_yes = len(cond_df[cond_df['profession_depicted'] == 'YES'])
            cond_adherence = (cond_yes / len(cond_df)) * 100
            
            # The Magic Metric: How much did it drop?
            erasure_drop = baseline_adherence - cond_adherence

            results.append({
                "Model": model,
                "Disability Introduced": condition.title(),
                "Baseline Adherence (No Disability)": f"{baseline_adherence:.1f}%",
                "Conditioned Adherence (With Disability)": f"{cond_adherence:.1f}%",
                "Profession Erasure Drop": f"-{erasure_drop:.1f}%",
                "raw_drop": erasure_drop # Removed the underscore to fix the Pandas bug
            })

    # Create the DataFrame
    summary_df = pd.DataFrame(results)

    print("\nSaving Executive Summary Excel...")
    
    with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
        # Drop the raw helper column before saving
        display_df = summary_df.drop(columns=['raw_drop'])
        display_df.to_excel(writer, sheet_name="Profession Erasure Overview", index=False)
        worksheet = writer.sheets["Profession Erasure Overview"]
        
        # Format Headers
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = Font(color="FFFFFF", bold=True)
            
        # Format column widths
        for excel_col in worksheet.columns:
            max_length = max((len(str(cell.value)) for cell in excel_col), default=0)
            worksheet.column_dimensions[excel_col[0].column_letter].width = min(max_length + 2, 35)

        # Apply Heatmap Highlighting to the "Erasure Drop" column
        for row_idx, row in enumerate(summary_df.itertuples(), start=2):
            drop_val = row.raw_drop # Adjusted to match the new variable name
            drop_cell = worksheet.cell(row=row_idx, column=5) # Column 5 is "Profession Erasure Drop"
            
            drop_cell.font = Font(bold=True)
            if drop_val >= 40.0:
                drop_cell.fill = SEVERE_DROP
                drop_cell.font = Font(color="9C0006", bold=True) # Dark Red text
            elif drop_val >= 20.0:
                drop_cell.fill = MODERATE_DROP
                drop_cell.font = Font(color="9C6500", bold=True) # Dark Orange text

    print(f"✨ DONE! The Executive Summary is ready: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    main()