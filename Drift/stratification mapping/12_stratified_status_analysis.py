import pandas as pd
import numpy as np
from scipy.stats import chi2_contingency
from pathlib import Path
from openpyxl.styles import PatternFill, Font

# ─── PATHS ──────────────────────────────────────────────────────────────────────
FLUX_ATTR_CSV = Path(r"C:\Study\Coding\Research\Drift\data_processed\01_flux_attributes_v2\attributes_v2.csv")
SD35_ATTR_CSV = Path(r"C:\Study\Coding\Research\Drift\data_processed\02_sd35_attributes_v2\attributes_v2.csv")

OUTPUT_DIR = Path(r"C:\Study\Coding\Research\Drift\stratification mapping")
OUTPUT_EXCEL = OUTPUT_DIR / "DCDD_Stratified_Status_Bias.xlsx"

# Automatically create the folder if it doesn't exist
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ─── BUCKETS (TAILORED TO YOUR LIST) ────────────────────────────────────────────
HIGH_STATUS = [
    'school teacher', 'university professor', 'doctor', 'accountant', 
    'investment analyst', 'software engineer', 'data scientist', 'journalist'
]

LOW_STATUS = [
    'hotel receptionist', 'waiter', 'bank teller', 'construction worker', 
    'electrician', 'factory worker', 'farmer', 'fisher', 'police officer', 
    'firefighter', 'chef', 'nurse'
]

HEADER_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
SIGNIFICANT_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

def calculate_cramers_v(confusion_matrix):
    chi2 = chi2_contingency(confusion_matrix)[0]
    n = confusion_matrix.sum().sum()
    r, k = confusion_matrix.shape
    if n == 0 or min(r, k) < 2: return 0.0
    return np.sqrt(chi2 / (n * (min(r, k) - 1)))

def assign_bucket(prompt_text):
    prompt_lower = str(prompt_text).lower()
    
    # Exclude non-professions explicitly
    if 'unemployed' in prompt_lower or 'retired' in prompt_lower:
        return "Exclude"
        
    for job in HIGH_STATUS:
        if job in prompt_lower: return "High-Status"
    for job in LOW_STATUS:
        if job in prompt_lower: return "Service/Manual"
    return "Other"

def main():
    print("[*] Loading attribute data for stratification...")
    
    try:
        flux_df = pd.read_csv(FLUX_ATTR_CSV)
        flux_df['model'] = 'FLUX.1 Dev'
        sd35_df = pd.read_csv(SD35_ATTR_CSV)
        sd35_df['model'] = 'SD 3.5 Large'
        df = pd.concat([flux_df, sd35_df], ignore_index=True)
    except FileNotFoundError as e:
        print(f"[!] Error: {e}")
        return

    # Create Intersectional Categories (Filtering out Ambiguous Noise!)
    df = df[~df['perceived_race'].astype(str).str.contains('Ambiguous', case=False, na=False)]
    df = df[~df['gender_presentation'].astype(str).str.contains('Ambiguous', case=False, na=False)]
    
    df['race_clean'] = df['perceived_race'].astype(str).str.replace(' presenting', '', regex=False).str.strip()
    df['gender_clean'] = df['gender_presentation'].astype(str).str.replace(' presenting', '', regex=False).str.strip()
    df['Intersection'] = df['race_clean'] + " - " + df['gender_clean']

    # Assign Buckets 
    text_cols = [c for c in df.columns if c.lower() in ['prompt', 'profession', 'text', 'job']]
    if not text_cols:
        print("[!] Could not find a text column to stratify by. (Make sure you have a 'prompt' column)")
        return
    
    target_col = text_cols[0]
    df['Status_Bucket'] = df[target_col].apply(assign_bucket)
    
    # Drop exclusions
    df = df[df['Status_Bucket'] != 'Exclude']
    
    results = []

    for bucket in ['High-Status', 'Service/Manual']:
        bucket_df = df[df['Status_Bucket'] == bucket]
        print(f"[*] Analyzing {bucket} ({len(bucket_df)} explicit demographic records)...")

        for model in ['FLUX.1 Dev', 'SD 3.5 Large']:
            model_df = bucket_df[bucket_df['model'] == model]
            if model_df.empty: continue

            neutral_df = model_df[model_df['image_type'] == 'neutral']
            condition_col = 'disability' if 'disability' in model_df.columns else 'condition'
            
            if condition_col not in model_df.columns:
                continue
                
            conditions = [d for d in model_df[condition_col].unique() if pd.notna(d) and str(d).strip() != 'neutral']

            for condition in conditions:
                cond_df = model_df[model_df[condition_col] == condition]
                if neutral_df.empty or cond_df.empty: continue
                    
                neutral_counts = neutral_df['Intersection'].value_counts()
                cond_counts = cond_df['Intersection'].value_counts()

                # --- THE FOOLPROOF SHIFT CALCULATION ---
                # Calculate percentages safely (avoiding divide-by-zero)
                neutral_pct = neutral_counts / max(neutral_counts.sum(), 1)
                cond_pct = cond_counts / max(cond_counts.sum(), 1)
                
                # Combine into a DataFrame to force safe alignment and fill missing categories with 0
                pct_df = pd.DataFrame({'neutral': neutral_pct, 'cond': cond_pct}).fillna(0)
                
                # Standard subtraction
                pct_df['delta'] = pct_df['cond'] - pct_df['neutral']
                
                if not pct_df.empty and pct_df['delta'].max() > 0:
                    true_shift_cat = pct_df['delta'].idxmax()
                    shift_magnitude = pct_df['delta'].max() * 100 
                else:
                    true_shift_cat = 'N/A'
                    shift_magnitude = 0.0

                all_categories = list(set(neutral_counts.index).union(set(cond_counts.index)))
                obs = [[neutral_counts.get(cat, 0), cond_counts.get(cat, 0)] for cat in all_categories]
                
                if len(obs) > 1 and sum(sum(row) for row in obs) > 0:
                    chi2, p_val, dof, expected = chi2_contingency(obs)
                    c_v = calculate_cramers_v(pd.DataFrame(obs))
                else:
                    p_val, c_v = 1.0, 0.0

                results.append({
                    "Status Bucket": bucket,
                    "Model": model,
                    "Disability Introduced": condition.title(),
                    "P-Value": p_val,
                    "Significant?": "YES" if p_val < 0.05 else "NO",
                    "Cramér's V": c_v,
                    "Demographic Shift": f"+{shift_magnitude:.1f}% towards {true_shift_cat}"
                })

    summary_df = pd.DataFrame(results)
    
    if not summary_df.empty:
        summary_df.sort_values(by=['Status Bucket', 'Model', 'P-Value'], inplace=True)

        with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
            summary_df.to_excel(writer, index=False)
            worksheet = writer.sheets["Sheet1"]
            for cell in worksheet[1]:
                cell.fill = HEADER_FILL
                cell.font = Font(color="FFFFFF", bold=True)
            for excel_col in worksheet.columns:
                worksheet.column_dimensions[excel_col[0].column_letter].width = 30
            for row_idx, row in enumerate(summary_df.itertuples(), start=2):
                if row._5 == "YES":  # Significant column
                    for col_idx in range(1, 8):
                        worksheet.cell(row=row_idx, column=col_idx).fill = SIGNIFICANT_FILL
                        worksheet.cell(row=row_idx, column=col_idx).font = Font(color="9C0006", bold=True)

        print(f"✨ DONE! Check {OUTPUT_EXCEL} for Stratified Status Bias.")
    else:
        print("[!] No results generated. Check if prompt/profession columns match.")

if __name__ == "__main__":
    main()